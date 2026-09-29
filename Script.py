import random
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# EMPLOYEE SALARY EXCEL GENERATOR
# ============================================================

OUTPUT_FILE = "Employee_Salary_Report.xlsx"
NUMBER_OF_EMPLOYEES = 200


# ------------------------------------------------------------
# Sample Data
# ------------------------------------------------------------

first_names = [
    "Aarav", "Aditi", "Akash", "Aman", "Ananya",
    "Arjun", "Ayush", "Bhavna", "Chetan", "Deepak",
    "Divya", "Gaurav", "Harish", "Isha", "Karan",
    "Kavya", "Krishna", "Manish", "Meera", "Mohit",
    "Neha", "Nikhil", "Nisha", "Pooja", "Pranav",
    "Priya", "Rahul", "Rakesh", "Riya", "Rohan",
    "Sakshi", "Sameer", "Sandeep", "Shivam", "Shreya",
    "Sneha", "Sonal", "Sumit", "Tanvi", "Varun",
    "Vikas", "Vivek", "Yash", "Zoya"
]

last_names = [
    "Sharma", "Verma", "Singh", "Patel", "Kumar",
    "Gupta", "Mehta", "Reddy", "Rao", "Iyer",
    "Nair", "Joshi", "Mishra", "Das", "Shah",
    "Chopra", "Malhotra", "Agarwal", "Bansal", "Kapoor"
]

departments = {
    "IT": [
        "Software Engineer",
        "Senior Software Engineer",
        "DevOps Engineer",
        "SRE",
        "Cloud Engineer"
    ],
    "HR": [
        "HR Executive",
        "HR Manager",
        "Talent Acquisition Specialist",
        "HR Business Partner"
    ],
    "Finance": [
        "Accountant",
        "Financial Analyst",
        "Finance Manager",
        "Senior Accountant"
    ],
    "Marketing": [
        "Marketing Executive",
        "Marketing Analyst",
        "Digital Marketing Specialist",
        "Marketing Manager"
    ],
    "Operations": [
        "Operations Executive",
        "Operations Analyst",
        "Operations Manager",
        "Process Associate"
    ],
    "Sales": [
        "Sales Executive",
        "Sales Manager",
        "Business Development Executive",
        "Sales Analyst"
    ],
    "Engineering": [
        "Engineer",
        "Senior Engineer",
        "Engineering Manager",
        "Technical Lead"
    ]
}


# ------------------------------------------------------------
# Salary ranges by department
# ------------------------------------------------------------

salary_ranges = {
    "IT": (50000, 150000),
    "HR": (35000, 100000),
    "Finance": (40000, 120000),
    "Marketing": (35000, 110000),
    "Operations": (30000, 90000),
    "Sales": (30000, 100000),
    "Engineering": (50000, 160000)
}


# ------------------------------------------------------------
# Generate Employee Data
# ------------------------------------------------------------

employees = []

for i in range(1, NUMBER_OF_EMPLOYEES + 1):

    employee_id = f"EMP{i:04d}"

    first_name = random.choice(first_names)
    last_name = random.choice(last_names)

    name = f"{first_name} {last_name}"

    department = random.choice(list(departments.keys()))

    designation = random.choice(departments[department])

    min_salary, max_salary = salary_ranges[department]

    basic_salary = random.randrange(
        min_salary,
        max_salary + 1,
        5000
    )

    # Salary Components
    hra = basic_salary * 0.40

    special_allowance = basic_salary * 0.20

    transport_allowance = 3000

    medical_allowance = 2000

    bonus = basic_salary * random.uniform(0.05, 0.15)

    # Gross Salary
    gross_salary = (
        basic_salary
        + hra
        + special_allowance
        + transport_allowance
        + medical_allowance
        + bonus
    )

    # Employee PF - 12% of Basic
    pf = basic_salary * 0.12

    # Professional Tax
    professional_tax = 200

    # Approximate TDS
    tds = gross_salary * random.uniform(0.05, 0.10)

    # Total deductions
    total_deductions = (
        pf
        + professional_tax
        + tds
    )

    # Net Salary
    net_salary = gross_salary - total_deductions

    # Employer PF
    employer_pf = basic_salary * 0.12

    # Annual CTC
    annual_ctc = (
        gross_salary
        + employer_pf
    ) * 12

    employees.append({
        "Employee ID": employee_id,
        "Employee Name": name,
        "Department": department,
        "Designation": designation,
        "Basic Salary": round(basic_salary, 2),
        "HRA": round(hra, 2),
        "Special Allowance": round(special_allowance, 2),
        "Transport Allowance": transport_allowance,
        "Medical Allowance": medical_allowance,
        "Bonus": round(bonus, 2),
        "Gross Salary": round(gross_salary, 2),
        "Employee PF": round(pf, 2),
        "Professional Tax": professional_tax,
        "TDS": round(tds, 2),
        "Total Deductions": round(total_deductions, 2),
        "Net Salary": round(net_salary, 2),
        "Employer PF": round(employer_pf, 2),
        "Annual CTC": round(annual_ctc, 2)
    })


# ------------------------------------------------------------
# Create DataFrame
# ------------------------------------------------------------

df = pd.DataFrame(employees)


# ------------------------------------------------------------
# Department Summary
# ------------------------------------------------------------

summary = (
    df.groupby("Department")
    .agg(
        Employees=("Employee ID", "count"),
        Total_Gross_Salary=("Gross Salary", "sum"),
        Average_Gross_Salary=("Gross Salary", "mean"),
        Total_Deductions=("Total Deductions", "sum"),
        Total_Net_Salary=("Net Salary", "sum"),
        Total_Annual_CTC=("Annual CTC", "sum")
    )
    .reset_index()
)


# ------------------------------------------------------------
# Overall Summary
# ------------------------------------------------------------

overall_summary = pd.DataFrame({
    "Metric": [
        "Total Employees",
        "Total Gross Salary",
        "Average Gross Salary",
        "Total Deductions",
        "Total Net Salary",
        "Total Annual CTC"
    ],

    "Value": [
        len(df),
        df["Gross Salary"].sum(),
        df["Gross Salary"].mean(),
        df["Total Deductions"].sum(),
        df["Net Salary"].sum(),
        df["Annual CTC"].sum()
    ]
})


# ------------------------------------------------------------
# Write Excel File
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Salary Details",
        index=False
    )

    summary.to_excel(
        writer,
        sheet_name="Department Summary",
        index=False
    )

    overall_summary.to_excel(
        writer,
        sheet_name="Overall Summary",
        index=False
    )


# ------------------------------------------------------------
# Excel Formatting
# ------------------------------------------------------------

workbook = load_workbook(OUTPUT_FILE)


# Colors
header_fill = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid"
)

summary_fill = PatternFill(
    start_color="D9EAF7",
    end_color="D9EAF7",
    fill_type="solid"
)

white_font = Font(
    color="FFFFFF",
    bold=True
)

bold_font = Font(
    bold=True
)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


# ------------------------------------------------------------
# Format Salary Details
# ------------------------------------------------------------

ws = workbook["Salary Details"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

for cell in ws[1]:

    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    cell.border = thin_border


# Currency columns
currency_columns = [
    "Basic Salary",
    "HRA",
    "Special Allowance",
    "Transport Allowance",
    "Medical Allowance",
    "Bonus",
    "Gross Salary",
    "Employee PF",
    "Professional Tax",
    "TDS",
    "Total Deductions",
    "Net Salary",
    "Employer PF",
    "Annual CTC"
]


header_map = {
    cell.value: cell.column
    for cell in ws[1]
}


for column_name in currency_columns:

    column_number = header_map[column_name]

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row=row,
            column=column_number
        ).number_format = '₹#,##0.00'


# Borders and alignment
for row in ws.iter_rows():

    for cell in row:

        cell.border = thin_border

        if cell.row > 1:
            cell.alignment = Alignment(
                vertical="center"
            )


# Column widths
for column in ws.columns:

    max_length = 0

    column_letter = get_column_letter(
        column[0].column
    )

    for cell in column:

        if cell.value is not None:

            max_length = max(
                max_length,
                len(str(cell.value))
            )

    ws.column_dimensions[
        column_letter
    ].width = min(max_length + 3, 30)


# ------------------------------------------------------------
# Format Department Summary
# ------------------------------------------------------------

ws = workbook["Department Summary"]

for cell in ws[1]:

    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(
        horizontal="center"
    )

    cell.border = thin_border


for row in ws.iter_rows():

    for cell in row:
        cell.border = thin_border


summary_currency_columns = [
    "Total_Gross_Salary",
    "Average_Gross_Salary",
    "Total_Deductions",
    "Total_Net_Salary",
    "Total_Annual_CTC"
]

header_map = {
    cell.value: cell.column
    for cell in ws[1]
}


for column_name in summary_currency_columns:

    column_number = header_map[column_name]

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row=row,
            column=column_number
        ).number_format = '₹#,##0.00'


for column in ws.columns:

    max_length = 0

    column_letter = get_column_letter(
        column[0].column
    )

    for cell in column:

        if cell.value is not None:

            max_length = max(
                max_length,
                len(str(cell.value))
            )

    ws.column_dimensions[
        column_letter
    ].width = min(max_length + 3, 30)


# ------------------------------------------------------------
# Format Overall Summary
# ------------------------------------------------------------

ws = workbook["Overall Summary"]

ws["A1"].fill = header_fill
ws["B1"].fill = header_fill

ws["A1"].font = white_font
ws["B1"].font = white_font

ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 25


for row in ws.iter_rows():

    for cell in row:
        cell.border = thin_border


for row in range(2, ws.max_row + 1):

    ws.cell(
        row=row,
        column=1
    ).font = bold_font

    metric = ws.cell(
        row=row,
        column=1
    ).value

    if metric != "Total Employees":

        ws.cell(
            row=row,
            column=2
        ).number_format = '₹#,##0.00'


# ------------------------------------------------------------
# Save Workbook
# ------------------------------------------------------------

workbook.save(OUTPUT_FILE)

print("=" * 60)
print("EMPLOYEE SALARY REPORT GENERATED SUCCESSFULLY")
print("=" * 60)
print(f"Total Employees : {NUMBER_OF_EMPLOYEES}")
print(f"Excel File      : {OUTPUT_FILE}")
print("=" * 60)